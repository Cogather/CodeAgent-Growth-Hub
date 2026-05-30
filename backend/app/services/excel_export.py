"""导入异常记录导出 Excel"""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font


def build_import_exceptions_excel(failures: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "导入异常"

    headers = ["工号", "姓名", "HR 返回部门路径", "异常原因"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for item in failures:
        dept_path = item.get("hr_dept_path") or []
        ws.append([
            item.get("emp_no", ""),
            item.get("name") or "",
            " / ".join(dept_path) if dept_path else "",
            item.get("reason", ""),
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_zone_import_exceptions_excel(failures: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "配置异常"

    headers = ["工号", "异常原因"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for item in failures:
        ws.append([item.get("emp_no", ""), item.get("reason", "")])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 48)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
