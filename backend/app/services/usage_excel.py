"""使用统计 Excel 解析与模板"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.services.dept_utils import HR_DEPT_STORE_LEVELS, hr_dept_visible_start

EMP_NO_HEADERS = {"工号", "员工工号", "emp_no", "emp no"}
USAGE_HEADERS = {"使用次数", "次数", "usage_count", "usage count", "使用量"}


@dataclass
class ParsedUsageRow:
    emp_no: str
    usage_count: int
    row_num: int


@dataclass
class UsageParseResult:
    rows: list[ParsedUsageRow]
    failures: list[dict[str, Any]]


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().lower().replace(" ", "_")


def _find_column(headers: list[str], candidates: set[str]) -> int | None:
    for idx, header in enumerate(headers):
        normalized = _normalize_header(header)
        if normalized in candidates or header.strip() in candidates:
            return idx
    return None


def parse_usage_excel(content: bytes) -> UsageParseResult:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        raise ValueError("Excel 文件为空")

    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        raise ValueError("Excel 缺少表头行")

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    emp_col = _find_column(headers, EMP_NO_HEADERS)
    usage_col = _find_column(headers, USAGE_HEADERS)

    if emp_col is None:
        raise ValueError("未找到「工号」列，请使用表头：工号、使用次数")
    if usage_col is None:
        raise ValueError("未找到「使用次数」列，请使用表头：工号、使用次数")

    parsed: list[ParsedUsageRow] = []
    failures: list[dict[str, Any]] = []
    seen_emp_nos: set[str] = set()

    for row_num, row in enumerate(rows_iter, start=2):
        if row is None:
            continue

        raw_emp = row[emp_col] if emp_col < len(row) else None
        raw_usage = row[usage_col] if usage_col < len(row) else None

        emp_no = str(raw_emp).strip() if raw_emp is not None else ""
        if not emp_no:
            if raw_usage is None or str(raw_usage).strip() == "":
                continue
            failures.append(
                {
                    "emp_no": "",
                    "usage_count": None,
                    "reason": f"第 {row_num} 行：工号为空",
                }
            )
            continue

        if emp_no in seen_emp_nos:
            failures.append(
                {
                    "emp_no": emp_no,
                    "usage_count": None,
                    "reason": f"第 {row_num} 行：工号重复",
                }
            )
            continue

        try:
            if raw_usage is None or str(raw_usage).strip() == "":
                raise ValueError("使用次数为空")
            if isinstance(raw_usage, bool):
                raise ValueError("使用次数格式无效")
            if isinstance(raw_usage, (int, float)):
                usage_count = int(raw_usage)
            else:
                usage_text = str(raw_usage).strip()
                if not re.fullmatch(r"-?\d+", usage_text):
                    raise ValueError("使用次数须为整数")
                usage_count = int(usage_text)
            if usage_count < 0:
                raise ValueError("使用次数不能为负数")
        except ValueError as exc:
            failures.append(
                {
                    "emp_no": emp_no,
                    "usage_count": None,
                    "reason": f"第 {row_num} 行：{exc}",
                }
            )
            continue

        seen_emp_nos.add(emp_no)
        parsed.append(ParsedUsageRow(emp_no=emp_no, usage_count=usage_count, row_num=row_num))

    wb.close()
    return UsageParseResult(rows=parsed, failures=failures)


def build_usage_template_excel() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "使用统计导入"
    ws.append(["工号", "使用次数"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append(["10001", 42])
    ws.append(["10002", 18])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 24)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_usage_import_failures_excel(failures: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "导入异常"

    headers = ["工号", "使用次数", "异常原因"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for item in failures:
        ws.append([
            item.get("emp_no", ""),
            item.get("usage_count") if item.get("usage_count") is not None else "",
            item.get("reason", ""),
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_zero_usage_excel(items: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "未使用人员"

    visible_start = hr_dept_visible_start()
    headers = ["工号", "姓名"]
    for level in range(visible_start, HR_DEPT_STORE_LEVELS + 1):
        headers.extend([f"{level}级部门", f"{level}级部门编码"])
    headers.append("使用次数")
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for item in items:
        row = [item.get("emp_no", "") or item.get("display_emp_no", ""), item.get("name", "")]
        for level in range(visible_start, HR_DEPT_STORE_LEVELS + 1):
            row.append(item.get(f"dept_l{level}_name") or "")
            row.append(item.get(f"dept_l{level}_code") or "")
        row.append(item.get("usage_count", 0))
        ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 36)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
